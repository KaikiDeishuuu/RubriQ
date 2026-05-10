from __future__ import annotations

from decimal import Decimal
from io import BytesIO
from types import SimpleNamespace
from typing import Any

from openpyxl import load_workbook

from fastapi.testclient import TestClient
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker

from app.api import deps
from app.db.base import Base
from app.main import app
from app.models import Answer, AnswerRubricResult, Exam, Question, RubricItem, Submission, SubmissionStatus
from app.services import export
from app.services.export import _answer_review_summary, _chunk_questions, _format_score_pair, _question_header, _review_status_label


def test_build_exam_results_pdf_returns_pdf_bytes(monkeypatch) -> None:
    def fake_results_data(_session: Any, exam_id: int) -> dict[str, Any]:
        return {
            "exam": SimpleNamespace(id=exam_id, title="Sample Exam"),
            "questions": [
                SimpleNamespace(question_no="1.1", max_score=4),
                SimpleNamespace(question_no="1.2", max_score=5),
            ],
            "ai_review_active": False,
            "ai_review_statuses": [],
            "rows": [
                {
                    "student_name": "Alice",
                    "student_id": "S001",
                    "status": "graded",
                    "total_score": 8.5,
                    "needs_human_review": False,
                    "question_scores": {"1.1": 4.0, "1.2": 4.5},
                },
                {
                    "student_name": "Bob",
                    "student_id": "S002",
                    "status": "needs_review",
                    "total_score": 3.0,
                    "needs_human_review": True,
                    "question_scores": {"1.1": 3.0},
                },
            ],
        }

    monkeypatch.setattr(export, "build_exam_results_data", fake_results_data)

    def fake_deductions_data(_session: Any, _exam_id: int, *, exam: Any = None) -> dict[str, Any]:
        return {
            "exam": exam,
            "questions": [],
            "per_question": [],
        }

    monkeypatch.setattr(export, "build_exam_deductions_data", fake_deductions_data)

    pdf_bytes = export.build_exam_results_pdf(None, 1)

    assert pdf_bytes.startswith(b"%PDF")
    assert len(pdf_bytes) > 1000


def test_chunk_questions_splits_long_result_tables() -> None:
    questions = [SimpleNamespace(question_no=str(index)) for index in range(1, 18)]

    chunks = _chunk_questions(questions, 8)

    assert [[question.question_no for question in chunk] for chunk in chunks] == [
        ["1", "2", "3", "4", "5", "6", "7", "8"],
        ["9", "10", "11", "12", "13", "14", "15", "16"],
        ["17"],
    ]


def test_teacher_finalized_is_not_exported_in_results_files(monkeypatch) -> None:
    def fake_results_data(_session: Any, exam_id: int) -> dict[str, Any]:
        return {
            "exam": SimpleNamespace(id=exam_id, title="Sample Exam"),
            "questions": [SimpleNamespace(question_no="1", max_score=5)],
            "ai_review_active": False,
            "ai_review_statuses": [],
            "rows": [
                {
                    "submission_id": 1,
                    "student_name": "Alice",
                    "student_id": "S001",
                    "status": "graded",
                    "total_score": 5.0,
                    "needs_human_review": False,
                    "teacher_finalized": True,
                    "question_scores": {"1": 5.0},
                }
            ],
        }

    monkeypatch.setattr(export, "build_exam_results_data", fake_results_data)

    csv_text = export.build_exam_results_csv(None, 1)
    workbook = load_workbook(BytesIO(export.build_exam_results_xlsx(None, 1)))
    xlsx_headers = [cell.value for cell in next(workbook.active.iter_rows(min_row=1, max_row=1))]

    assert "teacher_finalized" not in csv_text
    assert "终审" not in csv_text
    assert "平衡" not in csv_text
    assert "teacher_finalized" not in xlsx_headers
    assert "终审" not in xlsx_headers
    assert "平衡" not in xlsx_headers



def test_review_status_label_uses_clear_chinese_copy() -> None:
    assert _review_status_label(True) == "待复核"
    assert _review_status_label(False) == "已复核"



def test_question_score_pdf_copy_shows_score_and_max_score() -> None:
    question = SimpleNamespace(question_no="一", max_score=11)

    assert _question_header(question) == "一\n满分 11"
    assert _format_score_pair(8, question.max_score) == "8/11"
    assert _format_score_pair(8.5, question.max_score) == "8.5/11"
    assert _format_score_pair(None, question.max_score) == "-/11"


def test_answer_review_summary_describes_review_without_raw_response() -> None:
    answer = SimpleNamespace(
        review_triggers=["low_confidence", "score_delta"],
        review_score=Decimal("4"),
        score=Decimal("4"),
        max_score=Decimal("5"),
        review_decision="accepted_review",
    )

    summary = _answer_review_summary(answer)

    assert summary is not None
    assert "低置信度" in summary
    assert "快慢模型分差较大" in summary
    assert "4/5" in summary


def test_answer_review_summary_handles_failed_review() -> None:
    answer = SimpleNamespace(
        review_triggers=["score_variance"],
        review_score=None,
        score=Decimal("3"),
        max_score=Decimal("5"),
        review_decision="failed",
    )

    assert "强模型复审失败" in (_answer_review_summary(answer) or "")


def test_build_submission_review_pdf_returns_student_readable_pdf_bytes() -> None:
    engine = create_engine("sqlite:///:memory:", future=True)
    Base.metadata.create_all(engine)
    SessionLocal = sessionmaker(bind=engine, future=True)
    session = SessionLocal()
    try:
        exam = Exam(title="期中测验", total_score=Decimal("5"))
        session.add(exam)
        session.flush()
        question = Question(exam_id=exam.id, question_no="1", title="解释概念", max_score=Decimal("5"), order_index=1)
        session.add(question)
        session.flush()
        rubric_item = RubricItem(
            question_id=question.id,
            description="说明关键概念",
            max_score=Decimal("5"),
            order_index=1,
        )
        session.add(rubric_item)
        session.flush()
        submission = Submission(
            exam_id=exam.id,
            student_name="张三",
            student_id="S001",
            original_pdf_path="submissions/s001.pdf",
            status=SubmissionStatus.needs_review.value,
            total_score=Decimal("4"),
        )
        session.add(submission)
        session.flush()
        answer = Answer(
            submission_id=submission.id,
            question_id=question.id,
            extracted_answer="学生答案内容，包含较完整的概念说明。",
            score=Decimal("3.5"),
            max_score=Decimal("5"),
            ai_comment="整体理解正确，但论证略少。",
            missing_points=["缺少例子"],
            needs_human_review=False,
            teacher_override_score=Decimal("4"),
            teacher_comment="已复核，补充表述可给 4 分。",
            raw_ai_response="SHOULD_NOT_APPEAR",
            fast_score=Decimal("3.5"),
            review_score=Decimal("4"),
            review_triggers=["low_confidence"],
            review_decision="accepted_review",
        )
        session.add(answer)
        session.flush()
        session.add(
            AnswerRubricResult(
                answer_id=answer.id,
                rubric_item_id=rubric_item.id,
                awarded_score=Decimal("4"),
                evidence="答案提到了关键概念。",
                reason="关键定义准确，细节略少。",
            )
        )
        session.commit()

        pdf_bytes = export.build_submission_review_pdf(session, submission.id)
    finally:
        session.close()

    assert pdf_bytes.startswith(b"%PDF")
    assert len(pdf_bytes) > 1000



def test_export_exam_results_pdf_endpoint(monkeypatch) -> None:
    def fake_get_db():
        yield None

    monkeypatch.setattr("app.api.exams.build_exam_results_pdf", lambda _session, exam_id: b"%PDF-1.4 fake")
    app.dependency_overrides[deps.get_db] = fake_get_db
    try:
        response = TestClient(app).get("/api/exams/12/export.pdf")
    finally:
        app.dependency_overrides.clear()

    assert response.status_code == 200
    assert response.headers["content-type"] == "application/pdf"
    assert response.headers["content-disposition"] == 'attachment; filename="exam-12-results.pdf"'
    assert response.content == b"%PDF-1.4 fake"



def test_export_submission_review_pdf_endpoint(monkeypatch) -> None:
    def fake_get_db():
        yield None

    monkeypatch.setattr("app.api.submissions.build_submission_review_pdf", lambda _session, submission_id: b"%PDF-1.4 review")
    app.dependency_overrides[deps.get_db] = fake_get_db
    try:
        response = TestClient(app).get("/api/submissions/34/export.pdf")
    finally:
        app.dependency_overrides.clear()

    assert response.status_code == 200
    assert response.headers["content-type"] == "application/pdf"
    assert response.headers["content-disposition"] == 'attachment; filename="submission-34-review.pdf"'
    assert response.content == b"%PDF-1.4 review"


def test_build_exam_submissions_zip_packages_each_student(tmp_path) -> None:
    engine = create_engine("sqlite:///:memory:", future=True)
    Base.metadata.create_all(engine)
    SessionLocal = sessionmaker(bind=engine, future=True)
    session = SessionLocal()
    try:
        exam = Exam(title="集合测验", total_score=Decimal("5"))
        session.add(exam)
        session.flush()
        question = Question(exam_id=exam.id, question_no="1", title="说明集合", max_score=Decimal("5"), order_index=0)
        session.add(question)
        session.flush()
        rubric_item = RubricItem(
            question_id=question.id,
            description="给出定义",
            max_score=Decimal("5"),
            order_index=0,
        )
        session.add(rubric_item)
        session.flush()
        for student_id, student_name in [("S001", "张三"), ("S002", "李四"), ("BAD/ID", "王 五")]:
            submission = Submission(
                exam_id=exam.id,
                student_name=student_name,
                student_id=student_id,
                original_pdf_path=f"submissions/{student_id}.pdf",
                status=SubmissionStatus.graded.value,
                total_score=Decimal("4"),
            )
            session.add(submission)
            session.flush()
            answer = Answer(
                submission_id=submission.id,
                question_id=question.id,
                source_page=1,
                extracted_answer="一个简短的答案。",
                score=Decimal("4"),
                max_score=Decimal("5"),
                confidence="high",
                ai_comment="逻辑通顺，但缺少具体例子。",
                missing_points=["缺少例子"],
            )
            session.add(answer)
            session.flush()
            session.add(
                AnswerRubricResult(
                    answer_id=answer.id,
                    rubric_item_id=rubric_item.id,
                    awarded_score=Decimal("4"),
                    evidence="答案给出了集合的定义。",
                    reason="定义清晰，细节略少。",
                )
            )
        session.commit()

        zip_bytes, total, failures = export.build_exam_submissions_zip(session, exam.id)
    finally:
        session.close()

    assert total == 3
    assert failures == 0
    import io as _io
    import zipfile as _zip

    archive = _zip.ZipFile(_io.BytesIO(zip_bytes))
    names = sorted(archive.namelist())
    assert any("S001-张三" in name for name in names)
    assert any("S002-李四" in name for name in names)
    # Bad characters should be replaced rather than crashing.
    assert any("BAD_ID-王 五" in name or "BAD_ID-王_五" in name or "BAD_ID-王" in name for name in names)
    for name in names:
        with archive.open(name) as member:
            data = member.read()
            assert data.startswith(b"%PDF")


def test_build_exam_deductions_csv_lists_only_lost_points() -> None:
    engine = create_engine("sqlite:///:memory:", future=True)
    Base.metadata.create_all(engine)
    SessionLocal = sessionmaker(bind=engine, future=True)
    session = SessionLocal()
    try:
        exam = Exam(title="期中测验", total_score=Decimal("10"))
        session.add(exam)
        session.flush()
        question = Question(
            exam_id=exam.id,
            question_no="1.2",
            title="解释声阻抗",
            max_score=Decimal("10"),
            order_index=0,
        )
        session.add(question)
        session.flush()
        rubric_a = RubricItem(
            question_id=question.id,
            description="声阻抗与软组织相近",
            max_score=Decimal("6"),
            order_index=0,
        )
        rubric_b = RubricItem(
            question_id=question.id,
            description="减少反射",
            max_score=Decimal("4"),
            order_index=1,
        )
        session.add_all([rubric_a, rubric_b])
        session.flush()
        full_score_submission = Submission(
            exam_id=exam.id,
            student_name="李四",
            student_id="S002",
            original_pdf_path="submissions/s002.pdf",
            status=SubmissionStatus.graded.value,
            total_score=Decimal("10"),
        )
        session.add(full_score_submission)
        session.flush()
        session.add(
            Answer(
                submission_id=full_score_submission.id,
                question_id=question.id,
                extracted_answer="完整作答",
                score=Decimal("10"),
                max_score=Decimal("10"),
                ai_comment="完整。",
                missing_points=[],
            )
        )
        deducted_submission = Submission(
            exam_id=exam.id,
            student_name="张斯羽",
            student_id="S001",
            original_pdf_path="submissions/s001.pdf",
            status=SubmissionStatus.graded.value,
            total_score=Decimal("6"),
        )
        session.add(deducted_submission)
        session.flush()
        session.add(
            Answer(
                submission_id=deducted_submission.id,
                question_id=question.id,
                extracted_answer="只提到减少反射",
                score=Decimal("6"),
                max_score=Decimal("10"),
                ai_comment="只覆盖了减少反射。",
                missing_points=["未提及声阻抗"],
                teacher_override_score=None,
            )
        )
        overridden_submission = Submission(
            exam_id=exam.id,
            student_name="王 五",
            student_id="S003",
            original_pdf_path="submissions/s003.pdf",
            status=SubmissionStatus.graded.value,
            total_score=Decimal("8"),
        )
        session.add(overridden_submission)
        session.flush()
        session.add(
            Answer(
                submission_id=overridden_submission.id,
                question_id=question.id,
                extracted_answer="部分回答",
                score=Decimal("5"),
                max_score=Decimal("10"),
                ai_comment="AI 给 5 分。",
                missing_points=["不完整"],
                teacher_override_score=Decimal("8"),
                teacher_comment="教师认为应给 8 分。",
            )
        )
        session.commit()

        csv_text = export.build_exam_deductions_csv(session, exam.id)
        xlsx_bytes = export.build_exam_deductions_xlsx(session, exam.id)
    finally:
        session.close()

    rows = csv_text.splitlines()
    assert len(rows) == 3  # header + 2 deductions; full-score student excluded
    assert rows[0].lstrip("﻿").startswith("question_no,question_title")
    body = "\n".join(rows[1:])
    assert "S001" in body
    assert "S003" in body
    assert "S002" not in body
    assert "未提及声阻抗" in body
    assert "教师认为应给 8 分" in body
    workbook = load_workbook(BytesIO(xlsx_bytes))
    sheet = workbook.active
    assert sheet.cell(row=1, column=1).value == "question_no"
    student_ids = [sheet.cell(row=row_index, column=5).value for row_index in range(2, sheet.max_row + 1)]
    assert "S001" in student_ids
    assert "S003" in student_ids
    assert "S002" not in student_ids


def test_export_exam_deductions_endpoints(monkeypatch) -> None:
    def fake_get_db():
        yield None

    monkeypatch.setattr("app.api.exams._load_exam_or_404", lambda _session, exam_id: SimpleNamespace(id=exam_id))
    monkeypatch.setattr("app.api.exams.build_exam_deductions_csv", lambda _session, exam_id: "question_no\nfake")
    monkeypatch.setattr("app.api.exams.build_exam_deductions_xlsx", lambda _session, exam_id: b"PK fake xlsx")
    app.dependency_overrides[deps.get_db] = fake_get_db
    try:
        client = TestClient(app)
        csv_response = client.get("/api/exams/9/export-deductions.csv")
        xlsx_response = client.get("/api/exams/9/export-deductions.xlsx")
    finally:
        app.dependency_overrides.clear()

    assert csv_response.status_code == 200
    assert csv_response.headers["content-type"].startswith("text/csv")
    assert csv_response.headers["content-disposition"] == 'attachment; filename="exam-9-deductions.csv"'
    assert csv_response.text == "question_no\nfake"

    assert xlsx_response.status_code == 200
    assert xlsx_response.headers["content-type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    assert xlsx_response.headers["content-disposition"] == 'attachment; filename="exam-9-deductions.xlsx"'
    assert xlsx_response.content == b"PK fake xlsx"
