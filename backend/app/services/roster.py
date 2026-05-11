from __future__ import annotations

import csv
import io
import logging
import re
import time
from dataclasses import dataclass
from typing import Iterable

from sqlalchemy import delete, select
from sqlalchemy.orm import Session, selectinload

from app.core.config import settings
from app.models import Exam, ExamFile, RosterEntry
from app.schemas.ai import RosterParseResult
from app.schemas.exam import RosterEntryInput
from app.services.llm import call_structured_json
from app.services.ocr import build_ocr_reference_text, format_ocr_reference_text
from app.services.pdf import render_pdf_to_images
from app.storage.local import get_storage_service

logger = logging.getLogger(__name__)


_ROSTER_NAME_HEADERS = {"姓名", "学生姓名", "name", "student_name", "studentname", "student name"}
_ROSTER_ID_HEADERS = {"学号", "学籍号", "考号", "student_id", "studentid", "student id", "id"}


class RosterError(RuntimeError):
    pass


@dataclass(slots=True)
class ParsedRosterEntry:
    student_name: str | None
    student_id: str | None
    source: str = "manual"


@dataclass(slots=True)
class RosterMatch:
    entry: RosterEntry | None
    match_kind: str  # "id" / "name" / "none"


@dataclass(slots=True)
class RosterIndex:
    by_student_id: dict[str, RosterEntry]
    by_student_name: dict[str, RosterEntry]
    order_to_entry: dict[int, RosterEntry]
    entries: list[RosterEntry]

    def is_empty(self) -> bool:
        return not self.entries


def _normalize_id(value: str | None) -> str:
    if not value:
        return ""
    return re.sub(r"\s+", "", str(value)).casefold()


def _normalize_name(value: str | None) -> str:
    if not value:
        return ""
    return re.sub(r"\s+", "", str(value)).casefold()


def _clean_optional(value: str | None) -> str | None:
    if value is None:
        return None
    cleaned = str(value).strip()
    return cleaned or None


def load_roster_entries(session: Session, exam_id: int) -> list[RosterEntry]:
    return list(
        session.execute(
            select(RosterEntry)
            .where(RosterEntry.exam_id == exam_id)
            .order_by(RosterEntry.order_index.asc())
        ).scalars()
    )


def build_roster_index(session: Session, exam_id: int) -> RosterIndex:
    entries = load_roster_entries(session, exam_id)
    by_student_id: dict[str, RosterEntry] = {}
    by_student_name: dict[str, RosterEntry] = {}
    order_to_entry: dict[int, RosterEntry] = {}
    for entry in entries:
        order_to_entry[entry.order_index] = entry
        normalized_id = _normalize_id(entry.student_id)
        if normalized_id and normalized_id not in by_student_id:
            by_student_id[normalized_id] = entry
        normalized_name = _normalize_name(entry.student_name)
        if normalized_name and normalized_name not in by_student_name:
            by_student_name[normalized_name] = entry
    return RosterIndex(
        by_student_id=by_student_id,
        by_student_name=by_student_name,
        order_to_entry=order_to_entry,
        entries=entries,
    )


def match_roster_identity(
    index: RosterIndex,
    student_name: str | None,
    student_id: str | None,
) -> RosterMatch:
    if index.is_empty():
        return RosterMatch(entry=None, match_kind="none")
    normalized_id = _normalize_id(student_id)
    if normalized_id and normalized_id in index.by_student_id:
        return RosterMatch(entry=index.by_student_id[normalized_id], match_kind="id")
    normalized_name = _normalize_name(student_name)
    if normalized_name and normalized_name in index.by_student_name:
        return RosterMatch(entry=index.by_student_name[normalized_name], match_kind="name")
    return RosterMatch(entry=None, match_kind="none")


def replace_roster_entries(
    session: Session,
    exam_id: int,
    entries: Iterable[ParsedRosterEntry | RosterEntryInput],
    *,
    source: str = "manual",
) -> Exam:
    exam = session.get(Exam, exam_id)
    if exam is None:
        raise RosterError(f"Exam {exam_id} not found")
    session.execute(delete(RosterEntry).where(RosterEntry.exam_id == exam_id))
    session.flush()
    materialized: list[RosterEntry] = []
    order_index = 0
    for entry in entries:
        student_name = _clean_optional(getattr(entry, "student_name", None))
        student_id = _clean_optional(getattr(entry, "student_id", None))
        if student_name is None and student_id is None:
            continue
        roster_entry = RosterEntry(
            exam_id=exam_id,
            order_index=order_index,
            student_name=student_name,
            student_id=student_id,
            source=source,
        )
        session.add(roster_entry)
        materialized.append(roster_entry)
        order_index += 1
    if not materialized:
        raise RosterError("Roster must contain at least one entry with a name or id")
    exam.roster_status = "needs_review"
    exam.roster_error_message = None
    session.flush()
    return exam


def confirm_roster(session: Session, exam_id: int) -> Exam:
    exam = session.get(Exam, exam_id)
    if exam is None:
        raise RosterError(f"Exam {exam_id} not found")
    entries = load_roster_entries(session, exam_id)
    if not entries:
        raise RosterError("Cannot confirm an empty roster")
    if any(entry.student_name is None and entry.student_id is None for entry in entries):
        raise RosterError("Each roster entry must include a name or a student id")
    exam.roster_status = "confirmed"
    exam.roster_error_message = None
    session.commit()
    return exam


def clear_roster(session: Session, exam_id: int) -> Exam:
    exam = session.get(Exam, exam_id)
    if exam is None:
        raise RosterError(f"Exam {exam_id} not found")
    session.execute(delete(RosterEntry).where(RosterEntry.exam_id == exam_id))
    exam.roster_status = "not_uploaded"
    exam.roster_raw_ai_response = None
    exam.roster_error_message = None
    session.commit()
    return exam


def parse_roster_pdf(session: Session, exam_id: int, exam_file_id: int | None = None) -> Exam:
    exam = session.execute(
        select(Exam)
        .where(Exam.id == exam_id)
        .options(selectinload(Exam.files))
    ).scalar_one_or_none()
    if exam is None:
        raise RosterError(f"Exam {exam_id} not found")
    exam_file = _pick_roster_file(exam, exam_file_id)
    storage = get_storage_service()
    pdf_path = storage.path_for(exam_file.storage_path)
    if not pdf_path.exists():
        raise RosterError("Roster PDF no longer exists")
    rendered_pages = render_pdf_to_images(
        pdf_path,
        storage.path_for(f"rendered/exams/{exam.id}/roster/{exam_file.id}"),
        dpi=settings.render_dpi,
    )
    ocr_started_at = time.perf_counter()
    ocr_reference_text = build_ocr_reference_text(
        rendered_pages,
        "vision_roster",
        session=session,
        exam_id=exam.id,
    )
    if ocr_reference_text:
        logger.info(
            "Roster OCR reference built exam_id=%s pages=%s chars=%s duration_seconds=%.2f",
            exam_id,
            len(rendered_pages),
            len(ocr_reference_text),
            time.perf_counter() - ocr_started_at,
        )
    exam.roster_status = "parsing"
    exam.roster_error_message = None
    session.commit()
    try:
        completion = call_structured_json(
            model=settings.ai_vision_model,
            system_prompt_name="roster_extraction.system.md",
            user_prompt_name="roster_extraction.user.md",
            response_model=RosterParseResult,
            prompt_variables={
                "exam_title": exam.title,
                "ocr_reference_text": format_ocr_reference_text(ocr_reference_text),
            },
            image_paths=[page.image_path for page in rendered_pages],
            request_profile="vision",
            route_key="vision_roster",
        )
    except Exception as exc:  # noqa: BLE001 - parsing should surface a user-visible failure
        exam.roster_status = "needs_review"
        exam.roster_error_message = f"Roster parsing failed: {exc}"
        session.commit()
        raise RosterError(f"Roster parsing failed: {exc}") from exc

    parsed_entries = [
        ParsedRosterEntry(
            student_name=_clean_optional(student.student_name),
            student_id=_clean_optional(student.student_id),
            source="pdf",
        )
        for student in completion.data.students
    ]
    exam_file.parsed_json = RosterParseResult(students=completion.data.students).model_dump(mode="json")
    exam_file.raw_ai_response = completion.raw_text
    exam_file.error_message = None
    exam_file.page_count = len(rendered_pages)

    if not parsed_entries:
        exam.roster_status = "needs_review"
        exam.roster_raw_ai_response = completion.raw_text
        exam.roster_error_message = "AI 未在名单中识别到任何学生，请手动编辑"
        session.commit()
        return exam

    replace_roster_entries(session, exam_id, parsed_entries, source="pdf")
    exam.roster_raw_ai_response = completion.raw_text
    session.commit()
    return exam


def parse_roster_table(data: bytes, suffix: str) -> list[ParsedRosterEntry]:
    suffix_normalized = suffix.lower().lstrip(".")
    if suffix_normalized == "csv":
        return _parse_csv_bytes(data)
    if suffix_normalized in {"xlsx", "xlsm"}:
        return _parse_xlsx_bytes(data)
    raise RosterError(f"Unsupported roster table format: {suffix}")


def _parse_csv_bytes(data: bytes) -> list[ParsedRosterEntry]:
    text = _decode_text_bytes(data)
    return _rows_to_entries(_iter_csv_rows(text), source="csv")


def _parse_xlsx_bytes(data: bytes) -> list[ParsedRosterEntry]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover - declared in pyproject.toml
        raise RosterError("openpyxl is required to parse xlsx rosters") from exc
    workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    try:
        sheet = workbook.active
        rows: list[list[str]] = []
        for raw_row in sheet.iter_rows(values_only=True):
            rows.append([
                "" if cell is None else str(cell)
                for cell in raw_row
            ])
    finally:
        workbook.close()
    return _rows_to_entries(rows, source="xlsx")


def _decode_text_bytes(data: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-8", "gbk", "gb18030", "latin-1"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="ignore")


def _iter_csv_rows(text: str) -> list[list[str]]:
    sniffer = csv.Sniffer()
    sample = text[:4096]
    delimiter = ","
    try:
        dialect = sniffer.sniff(sample, delimiters=",\t;")
        delimiter = dialect.delimiter
    except csv.Error:
        pass
    reader = csv.reader(io.StringIO(text), delimiter=delimiter)
    return [[cell for cell in row] for row in reader]


def _rows_to_entries(rows: list[list[str]], *, source: str) -> list[ParsedRosterEntry]:
    if not rows:
        return []
    header_index = _detect_header_row(rows)
    if header_index is None:
        # Best effort: assume first column is name, second is id.
        body_rows = rows
        name_idx, id_idx = 0, 1
    else:
        header_row = [str(cell).strip().casefold() for cell in rows[header_index]]
        name_idx = _find_column(header_row, _ROSTER_NAME_HEADERS)
        id_idx = _find_column(header_row, _ROSTER_ID_HEADERS)
        if name_idx is None and id_idx is None:
            body_rows = rows
            name_idx, id_idx = 0, 1
        else:
            body_rows = rows[header_index + 1 :]
            if name_idx is None:
                name_idx = 0 if id_idx != 0 else 1
            if id_idx is None:
                id_idx = 1 if name_idx != 1 else 0
    entries: list[ParsedRosterEntry] = []
    for row in body_rows:
        if not any(str(cell).strip() for cell in row):
            continue
        student_name = _clean_optional(row[name_idx]) if name_idx is not None and name_idx < len(row) else None
        student_id = _clean_optional(row[id_idx]) if id_idx is not None and id_idx < len(row) else None
        if student_name is None and student_id is None:
            continue
        entries.append(
            ParsedRosterEntry(student_name=student_name, student_id=student_id, source=source)
        )
    return entries


def _detect_header_row(rows: list[list[str]]) -> int | None:
    for index, row in enumerate(rows[:5]):
        normalized = [str(cell).strip().casefold() for cell in row]
        if any(cell in _ROSTER_NAME_HEADERS for cell in normalized):
            return index
        if any(cell in _ROSTER_ID_HEADERS for cell in normalized):
            return index
    return None


def _find_column(header_row: list[str], candidates: set[str]) -> int | None:
    for index, cell in enumerate(header_row):
        if cell in candidates:
            return index
    return None


def _pick_roster_file(exam: Exam, exam_file_id: int | None) -> ExamFile:
    roster_files = [
        exam_file
        for exam_file in exam.files
        if (exam_file.file_type or "").startswith("roster_")
    ]
    if exam_file_id is not None:
        for exam_file in roster_files:
            if exam_file.id == exam_file_id:
                return exam_file
        raise RosterError(f"Exam file {exam_file_id} is not a roster file")
    if not roster_files:
        raise RosterError("No roster file uploaded for this exam")
    return sorted(roster_files, key=lambda item: item.created_at)[-1]
